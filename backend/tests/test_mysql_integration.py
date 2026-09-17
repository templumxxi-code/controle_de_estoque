import os
import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

mysql_url = os.getenv("MYSQL_TEST_DATABASE_URL")
if not mysql_url:
    pytest.skip("MYSQL_TEST_DATABASE_URL não configurada", allow_module_level=True)
os.environ["DATABASE_URL"] = mysql_url

from app.db.session import Base, SessionLocal, engine  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.main import app  # noqa: E402
from app.models import SaleItem, User  # noqa: E402


client = TestClient(app)


def clear_database() -> None:
    with engine.begin() as connection:
        connection.execute(text("SET FOREIGN_KEY_CHECKS=0"))
        for table in reversed(Base.metadata.sorted_tables):
            connection.execute(text(f"DELETE FROM `{table.name}`"))
        connection.execute(text("SET FOREIGN_KEY_CHECKS=1"))


def auth() -> dict[str, str]:
    with SessionLocal() as db:
        db.add(User(email="mysql-admin@test.local", password_hash=hash_password("senha-segura"), is_admin=True, role="admin", full_name="Administrador MySQL"))
        db.commit()
    response = client.post("/auth/login", json={"email": "mysql-admin@test.local", "password": "senha-segura"})
    assert response.status_code == 200, response.text
    return {"Authorization": f'Bearer {response.json()["access_token"]}'}


def test_mysql_end_to_end_and_rollback() -> None:
    clear_database()
    headers = auth()
    category = client.post("/categories", headers=headers, json={"name": "Suplementos"}).json()
    product = client.post("/products", headers=headers, json={"internal_code": "TEST-WHEY-001", "name": "Whey Teste", "category_id": category["id"], "purchase_cost": 50, "sale_price": 100, "desired_margin": 0.5, "minimum_stock": 3}).json()
    entry = client.post("/stock/entries", headers=headers, json={"product_id": product["id"], "quantity": 10, "unit_cost": 50})
    assert float(entry.json()["stock_quantity"]) == 10
    sale = client.post("/sales", headers=headers, json={"items": [{"product_id": product["id"], "quantity": 2, "unit_price": 100}]})
    assert sale.status_code == 201
    assert Decimal(str(sale.json()["total_amount"])) == Decimal("200")
    assert Decimal(str(sale.json()["total_cost"])) == Decimal("100")
    assert Decimal(str(sale.json()["gross_profit"])) == Decimal("100")
    products = client.get("/products", headers=headers).json()
    assert float(products[0]["stock_quantity"]) == 8
    changed = client.patch(f'/products/{product["id"]}', headers=headers, json={"internal_code": "TEST-WHEY-001", "name": "Whey Teste", "category_id": category["id"], "purchase_cost": 70, "sale_price": 100, "desired_margin": 0.3, "minimum_stock": 3}).json()
    assert float(changed["purchase_cost"]) == 70
    with SessionLocal() as db:
        item = db.query(SaleItem).filter(SaleItem.sale_id == sale.json()["id"]).one()
        assert Decimal(item.unit_cost) == Decimal("50.00")
        assert Decimal(item.gross_profit) == Decimal("100.00")
    client.post("/finance", headers=headers, json={"transaction_type": "expense", "category": "Operação", "description": "Despesa teste", "amount": 30})
    dashboard = client.get("/dashboard", headers=headers).json()
    assert Decimal(str(dashboard["revenue"])) == Decimal("200.00")
    assert Decimal(str(dashboard["gross_profit"])) == Decimal("100.00")
    assert Decimal(str(dashboard["net_result"])) == Decimal("70.00")
    blocked = client.post("/sales", headers=headers, json={"items": [{"product_id": product["id"], "quantity": 100}]})
    assert blocked.status_code == 409
    assert float(client.get("/products", headers=headers).json()[0]["stock_quantity"]) == 8
    assert len(client.get("/finance", headers=headers).json()) == 2
    for file_format in ("csv", "xlsx", "pdf"):
        exported = client.get(f"/reports/sales/export/{file_format}", headers=headers)
        assert exported.status_code == 200
        assert len(exported.content) > 20
        if file_format == "csv":
            assert "faturamento" in exported.content.decode("utf-8-sig")
        elif file_format == "xlsx":
            workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
            assert workbook.active.max_row >= 2
        else:
            assert exported.content.startswith(b"%PDF")


def test_mysql_attachment_and_import() -> None:
    headers = {"Authorization": f'Bearer {client.post("/auth/login", json={"email": "mysql-admin@test.local", "password": "senha-segura"}).json()["access_token"]}'}
    pdf = client.post("/attachments", headers=headers, files={"file": ("teste.pdf", b"%PDF-1.4 test", "application/pdf")})
    assert pdf.status_code == 201
    assert pdf.json()["size"] > 0
    png = client.post("/attachments", headers=headers, files={"file": ("teste.png", b"\x89PNG\r\n\x1a\n", "image/png")})
    assert png.status_code == 201
    assert png.json()["content_type"] == "image/png"
    invalid = client.post("/attachments", headers=headers, files={"file": ("teste.txt", b"invalid", "text/plain")})
    assert invalid.status_code == 415
    csv_content = "codigo;nome;categoria;unidade;quantidade;estoque_minimo;custo;preco;margem\nMYSQL-IMP-001;Importado MySQL;Suplementos;un;1;0;10;20;0,5\n".encode()
    preview = client.post("/imports/preview", headers=headers, files={"file": ("produtos.csv", csv_content, "text/csv")})
    assert preview.json()["valid"] == 1
    imported = client.post("/imports/products", headers=headers, files={"file": ("produtos.csv", csv_content, "text/csv")})
    assert imported.json()["imported"] == 1
    assert any(item["internal_code"] == "MYSQL-IMP-001" for item in client.get("/products", headers=headers).json())
