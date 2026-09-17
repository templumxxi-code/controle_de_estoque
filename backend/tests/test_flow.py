import os
import io

os.environ["DATABASE_URL"] = "sqlite:///./test_pumphouseup.db"
os.environ["SECRET_KEY"] = "test-secret"

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.security import hash_password
from app.db.session import Base, SessionLocal, engine
from app.main import app
from app.models import SaleItem, User

Base.metadata.drop_all(bind=engine)
Base.metadata.create_all(bind=engine)
with SessionLocal() as db:
    db.add(User(email="admin@pumphouseup.local", password_hash=hash_password("test-admin-password"), is_admin=True))
    db.commit()
client = TestClient(app)


def auth() -> dict[str, str]:
    response = client.post("/auth/login", json={"email": "admin@pumphouseup.local", "password": "test-admin-password"})
    return {"Authorization": f'Bearer {response.json()["access_token"]}'}


def test_acceptance_flow_and_stock_guard() -> None:
    headers = auth()
    product = client.post("/products", headers=headers, json={"internal_code": "WHEY-001", "name": "Whey", "purchase_cost": 50, "sale_price": 90, "desired_margin": 0.4, "minimum_stock": 2}).json()
    entry = client.post("/stock/entries", headers=headers, json={"product_id": product["id"], "quantity": 3, "unit_cost": 50}).json()
    assert float(entry["stock_quantity"]) == 3
    sale = client.post("/sales", headers=headers, json={"items": [{"product_id": product["id"], "quantity": 2}], "payment_method": "pix"})
    assert sale.status_code == 201
    assert float(sale.json()["gross_profit"]) == 80
    blocked = client.post("/sales", headers=headers, json={"items": [{"product_id": product["id"], "quantity": 2}]})
    assert blocked.status_code == 409
    dashboard = client.get("/dashboard", headers=headers).json()
    assert float(dashboard["revenue"]) == 180


def test_adjustment_snapshot_reports_and_exports() -> None:
    headers = auth()
    category = client.post("/categories", headers=headers, json={"name": "Proteínas"}).json()
    product = client.post("/products", headers=headers, json={"internal_code": "CREAT-001", "name": "Creatina", "category_id": category["id"], "purchase_cost": 50, "sale_price": 100, "desired_margin": 0.5, "minimum_stock": 2}).json()
    client.post("/stock/entries", headers=headers, json={"product_id": product["id"], "quantity": 10, "unit_cost": 50})
    sale = client.post("/sales", headers=headers, json={"items": [{"product_id": product["id"], "quantity": 2}]},).json()
    updated = client.patch(f'/products/{product["id"]}', headers=headers, json={"internal_code": "CREAT-001", "name": "Creatina", "category_id": category["id"], "purchase_cost": 70, "sale_price": 100, "desired_margin": 0.3, "minimum_stock": 2}).json()
    assert float(updated["purchase_cost"]) == 70
    with SessionLocal() as db:
        historical_item = db.query(SaleItem).filter(SaleItem.sale_id == sale["id"]).one()
        assert float(historical_item.unit_cost) == 50
        assert float(historical_item.gross_profit) == 100
    movements = client.get("/stock/movements", headers=headers, params={"product_id": product["id"]})
    assert movements.status_code == 200
    assert len(movements.json()) >= 2
    adjustment = client.post("/stock/adjustments", headers=headers, json={"product_id": product["id"], "adjustment_type": "increase", "quantity": 1, "reason": "Conferência física"})
    assert adjustment.status_code == 200
    assert float(adjustment.json()["stock_quantity"]) == 9
    report = client.get("/reports/sales", headers=headers)
    assert report.status_code == 200
    assert any(row["id"] == sale["id"] for row in report.json())
    for file_format in ("csv", "xlsx", "pdf"):
        exported = client.get(f"/reports/sales/export/{file_format}", headers=headers)
        assert exported.status_code == 200
        assert len(exported.content) > 10


def test_import_preview_and_invalid_authentication() -> None:
    assert client.get("/products").status_code == 401
    headers = auth()
    content = "codigo;nome;categoria;unidade;quantidade;estoque_minimo;custo;preco;margem\nIMP-001;Produto importado;Acessórios;un;2;1;10;20;0,50\n".encode("utf-8")
    preview = client.post("/imports/preview", headers=headers, files={"file": ("produtos.csv", content, "text/csv")})
    assert preview.status_code == 200
    assert preview.json()["valid"] == 1
    imported = client.post("/imports/products", headers=headers, files={"file": ("produtos.csv", content, "text/csv")})
    assert imported.status_code == 200
    assert imported.json()["imported"] == 1
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["codigo", "nome", "categoria", "unidade", "quantidade", "estoque_minimo", "custo", "preco", "margem"])
    sheet.append(["IMP-002", "Produto Excel", "Acessórios", "un", 1, 0, 12, 24, 0.5])
    xlsx = io.BytesIO(); workbook.save(xlsx)
    excel_preview = client.post("/imports/preview", headers=headers, files={"file": ("produtos.xlsx", xlsx.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert excel_preview.status_code == 200
    assert excel_preview.json()["valid"] == 1
    assert client.post("/auth/login", json={"email": "admin@test.local", "password": "errada"}).status_code == 401

def test_user_security_theme_and_audit() -> None:
    headers = auth()
    preference = client.put("/auth/preferences", headers=headers, json={"theme": "light"})
    assert preference.status_code == 200
    assert preference.json()["theme_preference"] == "light"
    created = client.post("/users", headers=headers, json={"email": "operator-v11@test.local", "password": "senha-segura", "full_name": "Operador V1.1", "role": "operator", "must_change_password": True})
    assert created.status_code == 201
    assert client.patch(f'/users/{created.json()["id"]}/status?active=false', headers=headers).status_code == 200
    assert client.post("/auth/login", json={"email": "operator-v11@test.local", "password": "senha-segura"}).status_code == 401
    assert client.patch(f'/users/{created.json()["id"]}/status?active=true', headers=headers).status_code == 200
    operator_login = client.post("/auth/login", json={"email": "operator-v11@test.local", "password": "senha-segura"})
    operator_headers = {"Authorization": f'Bearer {operator_login.json()["access_token"]}'}
    assert client.get("/users", headers=operator_headers).status_code == 403
    assert client.get("/audit-logs", headers=operator_headers).status_code == 403
    assert client.get("/settings", headers=operator_headers).status_code == 403
    assert client.get("/auth/me", headers=operator_headers).json()["must_change_password"] is True
    changed = client.post("/auth/change-password", headers=operator_headers, json={"current_password": "senha-segura", "new_password": "senha-nova-segura", "confirm_password": "senha-nova-segura"})
    assert changed.status_code == 200
    assert client.post("/auth/login", json={"email": "operator-v11@test.local", "password": "senha-nova-segura"}).status_code == 200
    actions = [item["action"] for item in client.get("/audit-logs", headers=headers).json()]
    assert {"LOGIN_SUCESSO", "USUARIO_CRIADO", "ALTERACAO_SENHA"}.issubset(actions)


def test_full_crud_catalog_finance_sale_and_excel_template() -> None:
    headers = auth()
    category = client.post("/categories", headers=headers, json={"name": "Linha QA"})
    assert category.status_code == 201
    category_id = category.json()["id"]
    updated_category = client.patch(f"/categories/{category_id}", headers=headers, json={"name": "Linha QA Atualizada"})
    assert updated_category.status_code == 200
    assert updated_category.json()["name"] == "Linha QA Atualizada"

    product_payload = {
        "internal_code": "QA-FULL-001",
        "name": "Produto Fluxo Completo",
        "category_id": category_id,
        "purchase_cost": 12,
        "sale_price": 25,
        "desired_margin": 0.5,
        "minimum_stock": 1,
    }
    product = client.post("/products", headers=headers, json=product_payload)
    assert product.status_code == 201
    product_id = product.json()["id"]
    assert client.post("/products", headers=headers, json=product_payload).status_code == 409
    product_payload["name"] = "Produto Fluxo Completo Atualizado"
    updated_product = client.patch(f"/products/{product_id}", headers=headers, json=product_payload)
    assert updated_product.status_code == 200
    assert updated_product.json()["name"] == "Produto Fluxo Completo Atualizado"

    entry = client.post("/stock/entries", headers=headers, json={"product_id": product_id, "quantity": 5, "unit_cost": 12, "reason": "Entrada QA"})
    assert entry.status_code == 200
    adjustment = client.post("/stock/adjustments", headers=headers, json={"product_id": product_id, "adjustment_type": "set", "quantity": 4, "reason": "Ajuste QA"})
    assert adjustment.status_code == 200
    assert float(adjustment.json()["stock_quantity"]) == 4

    expense = client.post("/finance", headers=headers, json={"transaction_type": "expense", "category": "QA", "description": "Despesa de teste", "amount": 30, "payment_method": "pix", "status": "paid"})
    assert expense.status_code == 201
    expense_id = expense.json()["id"]
    updated_expense = client.patch(f"/finance/{expense_id}", headers=headers, json={"transaction_type": "expense", "category": "QA Atualizada", "description": "Despesa atualizada", "amount": 35, "payment_method": "pix", "status": "paid"})
    assert updated_expense.status_code == 200
    assert float(updated_expense.json()["amount"]) == 35

    sale = client.post("/sales", headers=headers, json={"items": [{"product_id": product_id, "quantity": 2, "unit_price": 25}], "payment_method": "pix"})
    assert sale.status_code == 201
    assert float(sale.json()["total_amount"]) == 50
    assert float(client.get(f"/products?search=Atualizado", headers=headers).json()[0]["stock_quantity"]) == 2

    user = client.post("/users", headers=headers, json={"email": "qa-full-flow@test.local", "password": "senha-inicial", "full_name": "Usuário QA", "role": "operator"})
    assert user.status_code == 201
    user_id = user.json()["id"]
    updated_user = client.patch(f"/users/{user_id}", headers=headers, json={"full_name": "Usuário QA Atualizado", "job_title": "Operação"})
    assert updated_user.status_code == 200
    reset = client.post(f"/users/{user_id}/reset-password", headers=headers, json={"new_password": "senha-resetada"})
    assert reset.status_code == 200
    assert client.post("/auth/login", json={"email": "qa-full-flow@test.local", "password": "senha-resetada"}).status_code == 200

    template = client.get("/imports/template", headers=headers)
    assert template.status_code == 200
    workbook = load_workbook(io.BytesIO(template.content), read_only=True)
    assert list(workbook.active.values)[0] == ("codigo", "nome", "categoria", "unidade", "quantidade", "estoque_minimo", "custo", "preco", "margem")
