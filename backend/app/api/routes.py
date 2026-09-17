import csv
import io
import os
import uuid
import json
from datetime import datetime, timedelta, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.security import OAuth2PasswordBearer
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.security import create_access_token, decode_subject, hash_password, verify_password
from app.db.session import get_db
from app.models import Attachment, AuditLog, Category, FinancialTransaction, Product, Sale, SaleItem, StockMovement, SystemSetting, User
from app.schemas import AuditLogOut, CategoryCreate, CategoryOut, CategoryUpdate, FinancialCreate, FinancialUpdate, LoginRequest, PasswordChange, PasswordReset, ProductCreate, ProductOut, ProductUpdate, SaleCreate, StockAdjustment, StockEntry, StockMovementOut, ThemeUpdate, Token, UserCreate, UserOut, UserUpdate

router = APIRouter()
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/login")


def setting_value(db: Session, key: str, default: str = "") -> str:
    setting = db.scalar(select(SystemSetting).where(SystemSetting.key == key))
    return setting.value if setting else default


@router.get("/public/branding", response_model=dict)
def public_branding(db: Session = Depends(get_db)) -> dict[str, str | bool]:
    logo_path = setting_value(db, "company_logo_path")
    return {"company_name": setting_value(db, "company_name", "Pumphouseup Suplementos"), "logo_available": bool(logo_path and os.path.isfile(logo_path))}


@router.get("/public/logo")
def public_logo(db: Session = Depends(get_db)) -> FileResponse:
    logo_path = setting_value(db, "company_logo_path")
    if not logo_path or not os.path.isfile(logo_path):
        raise HTTPException(status_code=404, detail="Logo da empresa não configurada")
    return FileResponse(logo_path)


def current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)) -> User:
    email = decode_subject(token)
    user = db.scalar(select(User).where(User.email == email, User.is_active.is_(True))) if email else None
    if not user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Sessão inválida ou expirada")
    return user


def current_admin(user: User = Depends(current_user)) -> User:
    if not user.is_admin:
        raise HTTPException(status_code=403, detail="Acesso restrito ao administrador")
    return user


def audit(db: Session, action: str, description: str, user_id: int | None = None, entity: str | None = None, entity_id: int | None = None, result: str = "success") -> None:
    db.add(AuditLog(user_id=user_id, action=action, description=description, entity=entity, entity_id=entity_id, result=result))


@router.post("/auth/login", response_model=Token)
def login(payload: LoginRequest, db: Session = Depends(get_db)) -> Token:
    user = db.scalar(select(User).where(User.email == payload.email.lower().strip(), User.is_active.is_(True)))
    if not user or not verify_password(payload.password, user.password_hash):
        audit(db, "LOGIN_FALHA", f"Falha de login para {payload.email.lower().strip()}", result="failure")
        db.commit()
        raise HTTPException(status_code=401, detail="E-mail ou senha inválidos")
    user.last_login_at = datetime.now(timezone.utc).replace(tzinfo=None)
    audit(db, "LOGIN_SUCESSO", "Login realizado", user.id, "user", user.id)
    db.commit()
    return Token(access_token=create_access_token(user.email))


@router.get("/auth/me", response_model=UserOut)
def me(user: User = Depends(current_user)) -> User:
    return user


@router.put("/auth/preferences", response_model=UserOut)
def update_preferences(payload: ThemeUpdate, user: User = Depends(current_user), db: Session = Depends(get_db)) -> User:
    user.theme_preference = payload.theme
    db.commit()
    db.refresh(user)
    return user


@router.patch("/auth/profile", response_model=UserOut)
def update_profile(payload: UserUpdate, user: User = Depends(current_user), db: Session = Depends(get_db)) -> User:
    values = payload.model_dump(exclude_none=True)
    values.pop("role", None)
    values.pop("is_active", None)
    for field, value in values.items():
        setattr(user, field, value)
    db.commit()
    db.refresh(user)
    return user


@router.post("/auth/logout")
def logout(user: User = Depends(current_user), db: Session = Depends(get_db)) -> dict[str, str]:
    audit(db, "LOGOUT", "Sessão encerrada", user.id, "user", user.id)
    db.commit()
    return {"message": "Logout registrado"}


@router.post("/auth/change-password")
def change_password(payload: PasswordChange, user: User = Depends(current_user), db: Session = Depends(get_db)) -> dict[str, str]:
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status_code=400, detail="Senha atual inválida")
    if payload.new_password != payload.confirm_password:
        raise HTTPException(status_code=400, detail="A confirmação da nova senha não confere")
    if payload.current_password == payload.new_password:
        raise HTTPException(status_code=400, detail="A nova senha deve ser diferente da atual")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    audit(db, "ALTERACAO_SENHA", "Senha alterada", user.id, "user", user.id)
    db.commit()
    return {"message": "Senha alterada com sucesso"}


@router.get("/users", response_model=list[UserOut])
def list_users(_: User = Depends(current_admin), db: Session = Depends(get_db)) -> list[User]:
    return list(db.scalars(select(User).order_by(User.email)))


@router.post("/users", response_model=UserOut, status_code=201)
def create_user(payload: UserCreate, _: User = Depends(current_admin), db: Session = Depends(get_db)) -> User:
    user = User(email=payload.email.lower().strip(), password_hash=hash_password(payload.password), is_admin=payload.role == "admin", role=payload.role, full_name=payload.full_name, phone=payload.phone, job_title=payload.job_title, must_change_password=payload.must_change_password)
    db.add(user)
    try:
        db.commit()
        db.refresh(user)
        audit(db, "USUARIO_CRIADO", "Usuário criado", user.id, "user", user.id)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="E-mail já cadastrado") from exc
    return user


@router.patch("/users/{user_id}", response_model=UserOut)
def update_user(user_id: int, payload: UserUpdate, actor: User = Depends(current_admin), db: Session = Depends(get_db)) -> User:
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(user, field, value)
    if payload.role is not None:
        user.is_admin = payload.role == "admin"
    db.commit()
    audit(db, "USUARIO_EDITADO", "Usuário atualizado", actor.id, "user", user.id)
    db.commit()
    db.refresh(user)
    return user


@router.post("/users/{user_id}/reset-password")
def reset_user_password(user_id: int, payload: PasswordReset, actor: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict[str, str]:
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = True
    audit(db, "REDEFINICAO_SENHA", "Senha redefinida pelo administrador", actor.id, "user", user.id)
    db.commit()
    return {"message": "Senha redefinida com sucesso"}


@router.patch("/users/{user_id}/status", response_model=UserOut)
def set_user_status(user_id: int, active: bool = Query(...), actor: User = Depends(current_admin), db: Session = Depends(get_db)) -> User:
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    if user.id == actor.id and not active:
        raise HTTPException(status_code=400, detail="O administrador atual não pode se inativar")
    user.is_active = active
    audit(db, "USUARIO_ATIVADO" if active else "USUARIO_INATIVADO", "Status do usuário alterado", actor.id, "user", user.id)
    db.commit()
    return user


@router.get("/audit-logs", response_model=list[AuditLogOut])
def list_audit_logs(action: str | None = None, user_id: int | None = None, _: User = Depends(current_admin), db: Session = Depends(get_db)) -> list[AuditLog]:
    query = select(AuditLog).order_by(AuditLog.created_at.desc()).limit(200)
    if action:
        query = query.where(AuditLog.action == action)
    if user_id:
        query = query.where(AuditLog.user_id == user_id)
    return list(db.scalars(query))


@router.get("/categories", response_model=list[CategoryOut])
def list_categories(_: User = Depends(current_user), db: Session = Depends(get_db)) -> list[Category]:
    return list(db.scalars(select(Category).where(Category.is_active.is_(True)).order_by(Category.name)))


@router.post("/categories", response_model=CategoryOut, status_code=201)
def create_category(payload: CategoryCreate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Category:
    category = Category(name=payload.name.strip())
    db.add(category)
    try:
        db.commit()
        db.refresh(category)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Categoria já existente") from exc
    return category


@router.patch("/categories/{category_id}", response_model=CategoryOut)
def update_category(category_id: int, payload: CategoryUpdate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Category:
    category = db.get(Category, category_id)
    if not category:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    category.name = payload.name.strip()
    try:
        db.commit()
        db.refresh(category)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Categoria já existente") from exc
    return category


@router.get("/products", response_model=list[ProductOut])
def list_products(search: str | None = None, include_inactive: bool = False, _: User = Depends(current_user), db: Session = Depends(get_db)) -> list[Product]:
    query = select(Product).order_by(Product.name)
    if not include_inactive:
        query = query.where(Product.is_active.is_(True))
    if search:
        query = query.where(Product.name.ilike(f"%{search}%"))
    return list(db.scalars(query))


@router.post("/products", response_model=ProductOut, status_code=201)
def create_product(payload: ProductCreate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Product:
    product = Product(**payload.model_dump())
    db.add(product)
    try:
        db.commit()
        db.refresh(product)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código interno já cadastrado") from exc
    return product


@router.patch("/products/{product_id}", response_model=ProductOut)
def update_product(product_id: int, payload: ProductUpdate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Product:
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    for field, value in payload.model_dump().items():
        setattr(product, field, value)
    try:
        db.commit()
        db.refresh(product)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código interno já cadastrado") from exc
    return product


@router.patch("/products/{product_id}/status", response_model=ProductOut)
def set_product_status(product_id: int, active: bool = Query(...), _: User = Depends(current_user), db: Session = Depends(get_db)) -> Product:
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    product.is_active = active
    db.commit()
    db.refresh(product)
    return product


@router.post("/stock/entries", response_model=ProductOut)
def stock_entry(payload: StockEntry, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Product:
    product = db.get(Product, payload.product_id)
    if not product or not product.is_active:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    previous = Decimal(product.stock_quantity or 0)
    product.stock_quantity = previous + payload.quantity
    product.purchase_cost = payload.unit_cost
    db.add(StockMovement(product_id=product.id, movement_type="entry", quantity=payload.quantity, previous_quantity=previous, new_quantity=product.stock_quantity, unit_cost=payload.unit_cost, reason=payload.reason, document_number=payload.document_number))
    db.commit()
    db.refresh(product)
    return product


@router.post("/stock/adjustments", response_model=ProductOut)
def stock_adjustment(payload: StockAdjustment, _: User = Depends(current_user), db: Session = Depends(get_db)) -> Product:
    product = db.get(Product, payload.product_id)
    if not product or not product.is_active:
        raise HTTPException(status_code=404, detail="Produto não encontrado")
    previous = Decimal(product.stock_quantity or 0)
    if payload.adjustment_type == "set":
        new_quantity = payload.quantity
    elif payload.adjustment_type == "increase":
        new_quantity = previous + payload.quantity
    else:
        new_quantity = previous - payload.quantity
    if new_quantity < 0:
        raise HTTPException(status_code=409, detail="O ajuste não pode gerar estoque negativo")
    difference = new_quantity - previous
    product.stock_quantity = new_quantity
    db.add(StockMovement(product_id=product.id, movement_type="adjustment", quantity=difference, previous_quantity=previous, new_quantity=new_quantity, reason=payload.reason))
    db.commit()
    db.refresh(product)
    return product


@router.get("/stock/movements", response_model=list[StockMovementOut])
def stock_movements(product_id: int | None = None, movement_type: str | None = None, _: User = Depends(current_user), db: Session = Depends(get_db)) -> list[StockMovement]:
    query = select(StockMovement).order_by(StockMovement.created_at.desc())
    if product_id:
        query = query.where(StockMovement.product_id == product_id)
    if movement_type:
        query = query.where(StockMovement.movement_type == movement_type)
    return list(db.scalars(query))


@router.post("/sales", response_model=dict, status_code=201)
def create_sale(payload: SaleCreate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    sale = Sale(payment_method=payload.payment_method, notes=payload.notes)
    total = Decimal("0")
    total_cost = Decimal("0")
    pending_movements: list[StockMovement] = []
    for item in payload.items:
        product = db.get(Product, item.product_id)
        if not product or not product.is_active:
            raise HTTPException(status_code=404, detail=f"Produto {item.product_id} não encontrado")
        available = Decimal(product.stock_quantity or 0)
        if item.quantity > available:
            raise HTTPException(status_code=409, detail=f"Estoque insuficiente para {product.name}. Disponível: {available}")
        unit_price = item.unit_price if item.unit_price is not None else Decimal(product.sale_price or 0)
        unit_cost = Decimal(product.purchase_cost or 0)
        item_total = unit_price * item.quantity
        item_cost = unit_cost * item.quantity
        previous = available
        product.stock_quantity = available - item.quantity
        sale.items.append(SaleItem(product_id=product.id, quantity=item.quantity, unit_price=unit_price, unit_cost=unit_cost, total_amount=item_total, gross_profit=item_total - item_cost))
        pending_movements.append(StockMovement(product_id=product.id, movement_type="sale", quantity=-item.quantity, previous_quantity=previous, new_quantity=product.stock_quantity, unit_cost=unit_cost, reason="Venda"))
        total += item_total
        total_cost += item_cost
    sale.total_amount = total
    sale.total_cost = total_cost
    db.add(sale)
    db.flush()
    db.add_all(pending_movements)
    db.add(FinancialTransaction(transaction_type="income", category="Vendas", description=f"Venda #{sale.id}", amount=total, payment_method=payload.payment_method, status="received", source_type="sale", source_id=sale.id))
    db.commit()
    return {"id": sale.id, "total_amount": total, "total_cost": total_cost, "gross_profit": total - total_cost}


@router.post("/finance", response_model=dict, status_code=201)
def create_financial(payload: FinancialCreate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    transaction = FinancialTransaction(**payload.model_dump(exclude_none=True))
    db.add(transaction)
    db.commit()
    db.refresh(transaction)
    return {"id": transaction.id, "amount": transaction.amount}


@router.get("/finance", response_model=list[dict])
def list_finance(transaction_type: str | None = None, status_filter: str | None = None, _: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict]:
    query = select(FinancialTransaction).order_by(FinancialTransaction.transaction_date.desc())
    if transaction_type:
        query = query.where(FinancialTransaction.transaction_type == transaction_type)
    if status_filter:
        query = query.where(FinancialTransaction.status == status_filter)
    return [{"id": item.id, "transaction_type": item.transaction_type, "category": item.category, "description": item.description, "amount": item.amount, "payment_method": item.payment_method, "status": item.status, "transaction_date": item.transaction_date, "source_type": item.source_type, "source_id": item.source_id, "notes": item.notes} for item in db.scalars(query)]


@router.patch("/finance/{transaction_id}", response_model=dict)
def update_finance(transaction_id: int, payload: FinancialUpdate, _: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    transaction = db.get(FinancialTransaction, transaction_id)
    if not transaction:
        raise HTTPException(status_code=404, detail="Lançamento não encontrado")
    if transaction.source_type == "sale":
        raise HTTPException(status_code=409, detail="Receitas de venda são gerenciadas pelo módulo de vendas")
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(transaction, field, value)
    db.commit()
    return {"id": transaction.id, "amount": transaction.amount}


@router.get("/dashboard", response_model=dict)
def dashboard(_: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    sales = db.execute(select(func.coalesce(func.sum(Sale.total_amount), 0), func.coalesce(func.sum(Sale.total_amount - Sale.total_cost), 0)).where(Sale.created_at >= month_start)).one()
    expenses = db.scalar(select(func.coalesce(func.sum(FinancialTransaction.amount), 0)).where(FinancialTransaction.transaction_type == "expense", FinancialTransaction.transaction_date >= month_start)) or 0
    products = list(db.scalars(select(Product).where(Product.is_active.is_(True))))
    recent = list(db.scalars(select(StockMovement).order_by(StockMovement.created_at.desc()).limit(10)))
    return {"revenue": sales[0], "gross_profit": sales[1], "expenses": expenses, "net_result": sales[1] - expenses, "inventory_value": sum((Decimal(p.stock_quantity or 0) * Decimal(p.purchase_cost or 0) for p in products), Decimal("0")), "active_products": len(products), "low_stock": sum(1 for p in products if Decimal(p.stock_quantity or 0) <= Decimal(p.minimum_stock or 0)), "low_stock_products": [{"name": p.name, "stock": p.stock_quantity, "minimum": p.minimum_stock} for p in products if Decimal(p.stock_quantity or 0) <= Decimal(p.minimum_stock or 0)], "recent_movements": [{"product_id": m.product_id, "type": m.movement_type, "quantity": m.quantity, "created_at": m.created_at} for m in recent]}


def report_rows(report_type: str, start_date: datetime | None, end_date: datetime | None, db: Session) -> list[dict]:
    if report_type in {"stock", "low_stock"}:
        products = list(db.scalars(select(Product).order_by(Product.name)))
        if report_type == "low_stock":
            products = [p for p in products if Decimal(p.stock_quantity or 0) <= Decimal(p.minimum_stock or 0)]
        return [{"id": p.id, "codigo": p.internal_code, "produto": p.name, "estoque": p.stock_quantity, "minimo": p.minimum_stock, "custo": p.purchase_cost, "preco": p.sale_price, "valor_estoque": Decimal(p.stock_quantity or 0) * Decimal(p.purchase_cost or 0)} for p in products]
    if report_type == "movements":
        query = select(StockMovement).order_by(StockMovement.created_at.desc())
        if start_date: query = query.where(StockMovement.created_at >= start_date)
        if end_date: query = query.where(StockMovement.created_at <= end_date)
        return [{"id": m.id, "produto_id": m.product_id, "tipo": m.movement_type, "quantidade": m.quantity, "saldo_anterior": m.previous_quantity, "saldo_novo": m.new_quantity, "motivo": m.reason, "data": m.created_at} for m in db.scalars(query)]
    if report_type in {"sales", "revenue", "profit", "top_products"}:
        query = select(Sale).order_by(Sale.created_at.desc())
        if start_date: query = query.where(Sale.created_at >= start_date)
        if end_date: query = query.where(Sale.created_at <= end_date)
        sales = list(db.scalars(query))
        if report_type == "sales":
            return [{"id": s.id, "data": s.created_at, "faturamento": s.total_amount, "cmv": s.total_cost, "lucro_bruto": s.total_amount - s.total_cost, "pagamento": s.payment_method} for s in sales]
        if report_type == "revenue":
            return [{"id": s.id, "data": s.created_at, "faturamento": s.total_amount} for s in sales]
        if report_type == "profit":
            return [{"id": s.id, "data": s.created_at, "lucro_bruto": s.total_amount - s.total_cost} for s in sales]
        totals: dict[int, dict] = {}
        for sale in sales:
            for item in sale.items:
                row = totals.setdefault(item.product_id, {"produto_id": item.product_id, "quantidade": Decimal("0"), "faturamento": Decimal("0")})
                row["quantidade"] += item.quantity
                row["faturamento"] += item.total_amount
        return sorted(totals.values(), key=lambda row: row["quantidade"], reverse=True)
    query = select(FinancialTransaction).order_by(FinancialTransaction.transaction_date.desc())
    if start_date: query = query.where(FinancialTransaction.transaction_date >= start_date)
    if end_date: query = query.where(FinancialTransaction.transaction_date <= end_date)
    transactions = list(db.scalars(query))
    if report_type == "income": transactions = [item for item in transactions if item.transaction_type == "income"]
    if report_type == "expenses": transactions = [item for item in transactions if item.transaction_type == "expense"]
    return [{"id": item.id, "tipo": item.transaction_type, "categoria": item.category, "descricao": item.description, "valor": item.amount, "status": item.status, "data": item.transaction_date, "origem": item.source_type} for item in transactions]


@router.get("/reports/{report_type}", response_model=list[dict])
def report(report_type: str, start_date: datetime | None = None, end_date: datetime | None = None, _: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict]:
    allowed = {"stock", "low_stock", "movements", "sales", "revenue", "profit", "income", "expenses", "financial", "top_products"}
    if report_type not in allowed:
        raise HTTPException(status_code=400, detail="Tipo de relatório inválido")
    return report_rows(report_type, start_date, end_date, db)


@router.get("/reports/{report_type}/export/{file_format}")
def export_report(report_type: str, file_format: str, start_date: datetime | None = None, end_date: datetime | None = None, _: User = Depends(current_user), db: Session = Depends(get_db)) -> StreamingResponse:
    rows = report_rows(report_type, start_date, end_date, db)
    if not rows:
        rows = [{"resultado": "sem dados"}]
    columns = list(rows[0])
    values = [[str(row.get(column, "")) for column in columns] for row in rows]
    filename = f"pumphouseup-{report_type}"
    if file_format == "csv":
        output = io.StringIO()
        writer = csv.writer(output, delimiter=";")
        writer.writerow(columns); writer.writerows(values)
        return StreamingResponse(iter([output.getvalue().encode("utf-8-sig")]), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}.csv"})
    if file_format == "xlsx":
        workbook = Workbook(); sheet = workbook.active
        sheet.append(columns)
        for row in values: sheet.append(row)
        output = io.BytesIO(); workbook.save(output); output.seek(0)
        return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"})
    if file_format == "pdf":
        output = io.BytesIO(); document = SimpleDocTemplate(output, pagesize=landscape(A4), title=filename)
        table = Table([columns] + values, repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111214")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 7)]))
        document.build([table]); output.seek(0)
        return StreamingResponse(output, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}.pdf"})
    raise HTTPException(status_code=400, detail="Formato de exportação inválido")


def parse_import(content: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text), delimiter=";"))
    if filename.lower().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.values)
        headers = [str(value or "").strip() for value in rows[0]] if rows else []
        return [dict(zip(headers, row)) for row in rows[1:]]
    raise HTTPException(status_code=400, detail="Use arquivo .csv ou .xlsx")


def validate_import(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    required = {"codigo", "nome", "categoria", "unidade", "quantidade", "estoque_minimo", "custo", "preco", "margem"}
    valid: list[dict] = []
    errors: list[dict] = []
    if rows and not required.issubset({str(key).strip().lower() for key in rows[0]}):
        missing = sorted(required - {str(key).strip().lower() for key in rows[0]})
        return [], [{"linha": 1, "campo": "colunas", "motivo": f"Colunas ausentes: {', '.join(missing)}"}]
    for line_number, original in enumerate(rows, start=2):
        row = {str(key).strip().lower(): value for key, value in original.items()}
        for field in required:
            if row.get(field) in (None, ""):
                errors.append({"linha": line_number, "campo": field, "motivo": "Campo obrigatório"})
        try:
            for field in ("quantidade", "estoque_minimo", "custo", "preco", "margem"):
                Decimal(str(row.get(field, "0")).replace(",", "."))
            if Decimal(str(row.get("margem", "0")).replace(",", ".")) >= 1:
                raise ValueError("margem deve ser menor que 100%")
            valid.append(row)
        except (ValueError, TypeError, ArithmeticError) as error:
            errors.append({"linha": line_number, "campo": "valores", "motivo": str(error)})
    return valid if not errors else [], errors


@router.post("/imports/preview", response_model=dict)
def import_preview(file: UploadFile = File(...), _: User = Depends(current_user)) -> dict:
    rows = parse_import(file.file.read(), file.filename or "")
    valid, errors = validate_import(rows)
    return {"filename": file.filename, "total": len(rows), "valid": len(valid), "errors": errors, "preview": rows[:10]}


@router.post("/imports/products", response_model=dict)
def import_products(file: UploadFile = File(...), _: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    rows = parse_import(file.file.read(), file.filename or "")
    valid, errors = validate_import(rows)
    if errors:
        return {"imported": 0, "errors": errors}
    imported = 0
    for row in valid:
        category = db.scalar(select(Category).where(Category.name == str(row["categoria"]).strip()))
        if not category:
            category = Category(name=str(row["categoria"]).strip()); db.add(category); db.flush()
        product = Product(internal_code=str(row["codigo"]).strip(), name=str(row["nome"]).strip(), category_id=category.id, unit=str(row["unidade"]).strip(), stock_quantity=Decimal(str(row["quantidade"]).replace(",", ".")), minimum_stock=Decimal(str(row["estoque_minimo"]).replace(",", ".")), purchase_cost=Decimal(str(row["custo"]).replace(",", ".")), sale_price=Decimal(str(row["preco"]).replace(",", ".")), desired_margin=Decimal(str(row["margem"]).replace(",", ".")))
        db.add(product); imported += 1
    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Código interno duplicado na importação") from exc
    return {"imported": imported, "errors": []}


@router.get("/imports/template")
def import_template(_: User = Depends(current_user)) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Produtos"
    sheet.append(["codigo", "nome", "categoria", "unidade", "quantidade", "estoque_minimo", "custo", "preco", "margem"])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=template-produtos.xlsx"})


@router.post("/attachments", status_code=201)
def upload_attachment(file: UploadFile = File(...), entity_type: str | None = Form(default=None), entity_id: int | None = Form(default=None), _: User = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    upload_dir = settings.upload_dir
    os.makedirs(upload_dir, exist_ok=True)
    allowed = {"application/pdf", "image/jpeg", "image/png", "image/webp"}
    if file.content_type not in allowed:
        raise HTTPException(status_code=415, detail="Tipo de arquivo não permitido")
    safe_name = os.path.basename(file.filename or "arquivo")
    stored_name = f"{uuid.uuid4().hex}-{safe_name}"
    target = os.path.join(upload_dir, stored_name)
    content = file.file.read()
    with open(target, "wb") as output:
        output.write(content)
    attachment = Attachment(original_name=safe_name, stored_path=target, content_type=file.content_type or "application/octet-stream", file_size=len(content), entity_type=entity_type, entity_id=entity_id)
    db.add(attachment); db.commit(); db.refresh(attachment)
    return {"id": attachment.id, "name": safe_name, "path": target, "content_type": file.content_type, "size": len(content)}


@router.get("/settings", response_model=dict)
def get_settings(_: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict:
    return {item.key: item.value for item in db.scalars(select(SystemSetting).order_by(SystemSetting.key))}


@router.post("/settings/logo", response_model=dict)
def upload_company_logo(file: UploadFile = File(...), user: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict[str, str]:
    allowed = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}
    extension = allowed.get(file.content_type or "")
    if not extension:
        raise HTTPException(status_code=415, detail="Use uma imagem PNG, JPG ou WEBP")
    os.makedirs(settings.upload_dir, exist_ok=True)
    target = os.path.abspath(os.path.join(settings.upload_dir, f"company-logo{extension}"))
    content = file.file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="A logo deve ter no máximo 5 MB")
    with open(target, "wb") as output:
        output.write(content)
    setting = db.scalar(select(SystemSetting).where(SystemSetting.key == "company_logo_path"))
    if setting:
        setting.value = target
    else:
        db.add(SystemSetting(key="company_logo_path", value=target))
    audit(db, "CONFIGURACAO_ALTERADA", "Logo da empresa atualizada", user.id, "settings")
    db.commit()
    return {"message": "Logo atualizada com sucesso", "logo_url": "/public/logo"}


@router.put("/settings", response_model=dict)
def update_settings(values: dict[str, str], _: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict:
    for key, value in values.items():
        setting = db.scalar(select(SystemSetting).where(SystemSetting.key == key))
        if setting:
            setting.value = value
        else:
            db.add(SystemSetting(key=key, value=value))
    db.commit()
    return get_settings(_, db)
